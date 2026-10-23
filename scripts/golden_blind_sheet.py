#!/usr/bin/env python3
"""文件層級 golden set 盲標表與 qrels 匯出（規格 Q4）。

build：由 docs/golden-set/2026-10-07-baroque-candidates.xlsx（15 題、371 筆）產生盲標表
  - 隱藏「來源」「被哪些策略找到（名次）」，每題內以固定種子打亂順序（不洩漏系統排名）
  - 標籤改 0／1／2：2 直接回答查詢、1 含可用佐證、0 只共用關鍵字（Recall／MRR 以 ≥1 為相關）
  - 每題一列「相關性說明」（2–3 句，先寫再標）；對照（盲標序 → doc_id、external_id）另存 _key.json
qrels：讀回標好的盲標表，輸出 TREC qrels（題號 0 doc_id 等級），供 ranx／ir_measures 計算

用法（py311_venv）：
    python scripts/golden_blind_sheet.py build <候選.xlsx> <盲標.xlsx>
    python scripts/golden_blind_sheet.py qrels <盲標.xlsx> <qrels.txt>
"""

import json
import random
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font

SEED = 20261014


def build(src, dst):
    wb = openpyxl.load_workbook(src, read_only=True)
    questions = list(wb["題目"].iter_rows(min_row=2, values_only=True))
    cands = list(wb["候選"].iter_rows(min_row=2, values_only=True))
    rng = random.Random(SEED)  # noqa: S311 -- 抽樣用固定種子，非安全用途
    out = openpyxl.Workbook()
    qs = out.active
    qs.title = "題目與相關性說明"
    qs.append(["題號", "題型", "題目", "候選數", "相關性說明（請先寫 2–3 句：怎樣算 2、怎樣算 1）"])
    for q in questions:
        qs.append([q[0], q[1], q[2], q[3], ""])
    ws = out.create_sheet("盲標")
    ws.append(["題號", "盲標序", "標題", "藝術家", "年代", "內容摘要", "等級（2／1／0）", "備註"])
    key = {}
    for no in sorted({c[0] for c in cands if c[0] is not None}):
        rows = [c for c in cands if c[0] == no]
        rng.shuffle(rows)
        for i, c in enumerate(rows, 1):
            ws.append([no, i, c[3], c[4], c[5], c[7], "", ""])
            key[f"{no}-{i}"] = {"doc_id": c[10], "external_id": c[9]}
    for sheet in (qs, ws):
        for c in sheet[1]:
            c.font = Font(bold=True)
        for row in sheet.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDEFGH", (6, 8, 40, 22, 10, 90, 14, 20), strict=True):
        ws.column_dimensions[col].width = w
    qs.column_dimensions["C"].width = 60
    qs.column_dimensions["E"].width = 70
    guide = out.create_sheet("標註說明")
    for line in [
        "等級 2：文件直接回答查詢（例：問《夜巡》媒材，文件寫明 Oil on canvas）。",
        "等級 1：文件含可用佐證但不直接回答（例：同畫家同時期作品，可支撐脈絡）。",
        "等級 0：只共用關鍵字，對回答沒有幫助。",
        "先在「題目與相關性說明」寫該題的判準，再標該題全部候選；一次標完一題，每段 45–60 分鐘。",
        "盲標序已打亂，與系統排名無關；請勿參考任何系統輸出或 LLM。",
        "標完約 2 週後會再抽 10%（37–40 筆）重標，計算一致性（目標 κ ≥ 0.7）。",
    ]:
        guide.append([line])
    guide.column_dimensions["A"].width = 110
    out.save(dst)
    Path(dst).with_name(Path(dst).stem + "_key.json").write_text(
        json.dumps(key, ensure_ascii=False, indent=1), encoding="utf-8"
    )
    print(f"盲標表 {dst}：{len(questions)} 題、{len(cands)} 筆；對照 {Path(dst).stem}_key.json")


def qrels(src, dst):
    key = json.loads(Path(src).with_name(Path(src).stem + "_key.json").read_text(encoding="utf-8"))
    wb = openpyxl.load_workbook(src, read_only=True)
    lines, missing = [], 0
    for no, seq, *_rest in wb["盲標"].iter_rows(min_row=2, values_only=True):
        grade = _rest[4]
        if grade in (None, ""):
            missing += 1
            continue
        lines.append(f"{no} 0 {key[f'{no}-{seq}']['doc_id']} {int(grade)}")
    Path(dst).write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"qrels {dst}：{len(lines)} 筆；未標 {missing} 筆")


if __name__ == "__main__":
    {"build": build, "qrels": qrels}[sys.argv[1]](sys.argv[2], sys.argv[3])
