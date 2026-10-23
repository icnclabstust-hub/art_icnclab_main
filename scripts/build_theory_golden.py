#!/usr/bin/env python3
"""理論題黃金測試集：把「15題」改寫版 docx（每題一檔）解析成結構化 JSON。

輸入：Q01_rewrite.docx … Q15_rewrite.docx（2026-10-13 依評分鍵查核與原典補查重新整理）
每個主張是正文段落，句末標【可信度 A–D｜忠實度 F0–F3｜文獻】：
- 可信度：A＝一手史料可直接證明、B＝權威研究支持、C＝推論或類比、D＝待核（可為「A／B」）
- 忠實度：F3＝已開原始頁或原典核對、F2＝書目真實但內文未逐頁核對、F1＝部分支撐、—＝本文提出
評估只用主要等級（第一個字母）為 A／B 的主張（與 2026-10-08 版的 grade in "AB" 相容）；
複合等級如「A（文本）／C（詮釋）」記 grade="A"、grades=["A","C"]、grade_raw 保留原文。
說明文字（原 prompt、改寫版說明）不算主張；正文內容表格另存 tables，不算主張。

輸出格式與 docs/golden-set/2026-10-08-theory15-claims.json 相容（no／title／question／summary／
claims[{grade, text}]），另加 fidelity、refs、section、references、pending、revisions。

用法（py311_venv，需 python-docx）：
    python scripts/build_theory_golden.py <docx 目錄> <輸出.json>
"""

import json
import re
import sys
from pathlib import Path

import docx

TAG = re.compile(r"\s*【可信度\s*([^｜】]+?)\s*｜\s*忠實度\s*([^｜】]*?)\s*｜\s*([^】]*)】\s*$")


def grade_letters(raw):
    """「A（文本）／C（詮釋）」→ ["A", "C"]；括號內的字母（如「B（物理常數 A）」的 A）不算。"""
    outside = re.sub(r"（[^）]*）|\([^)]*\)", "", raw)
    return re.findall(r"[A-D]", outside)


def style(p):
    return p.style.name if p.style is not None else ""


def parse(path):
    d = docx.Document(path)
    q = {
        "no": int(re.search(r"Q(\d+)", path.name).group(1)),
        "title": "",
        "question": "",
        "summary": "",
        "claims": [],
        "references": [],
        "pending": [],
        "revisions": [],
        "source_file": path.name,
    }
    part, section = None, ""
    for p in d.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        st = style(p)
        if st == "Heading 1":
            q["question"] = text
            continue
        if st == "Heading 2":
            part = text
            continue
        if st == "Heading 3":
            section = re.sub(r"^[一二三四五六七八九十]+、", "", text)
            continue
        if text.startswith("原 prompt："):
            q["title"] = text.removeprefix("原 prompt：").strip()
            continue
        if part == "執行摘要":
            q["summary"] = (q["summary"] + text).strip()
        elif part and part.startswith("修訂說明"):
            q["revisions"].append(re.sub(r"^\d+\.\s*", "", text))
        elif part and part.startswith("待查"):
            if re.match(r"^\d+\.", text):
                q["pending"].append(re.sub(r"^\d+\.\s*", "", text))
        elif part == "正文":
            m = TAG.search(text)
            if not m:
                continue
            body = text[: m.start()].lstrip("・").strip()
            raw, fidelity, refs = m.group(1).strip(), m.group(2).strip(), m.group(3).strip()
            letters = grade_letters(raw)
            if not letters:
                raise ValueError(f"{path.name}：無法判讀可信度「{raw}」")
            q["claims"].append(
                {
                    "grade": letters[0],
                    "grades": letters,
                    "grade_raw": raw,
                    "text": body,
                    "fidelity": fidelity,
                    "refs": refs,
                    "section": section,
                }
            )
    q["tables"] = []
    for t in d.tables:
        rows = [[c.text.strip() for c in row.cells] for row in t.rows]
        if rows and rows[0] and rows[0][0] == "編號":  # 參考文獻表
            q["references"] += [
                {"id": r[0], "citation": r[1], "status": r[2]}
                for r in rows[1:]
                if len(r) >= 3 and r[0]
            ]
        else:  # 正文內容表格（節點表、比較表），保留供參，不算主張
            q["tables"].append(rows)
    return q


def main():
    src, out = Path(sys.argv[1]), Path(sys.argv[2])
    items = [parse(p) for p in sorted(src.glob("Q*_rewrite.docx"))]
    out.write_text(json.dumps(items, ensure_ascii=False, indent=1), encoding="utf-8")
    for q in items:
        grades = [c["grade"] for c in q["claims"]]
        ab = sum(1 for g in grades if g in "AB")
        print(
            f"Q{q['no']:02} 主張 {len(grades):2}（A/B {ab:2}）文獻 {len(q['references']):2} "
            f"表格 {len(q['tables'])} 待查 {len(q['pending'])} | {q['question'][:36]}"
        )


if __name__ == "__main__":
    main()
