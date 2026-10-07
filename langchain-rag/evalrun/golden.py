"""黃金集匯入：讀人工標註的候選表 Excel，寫入 ``eval.golden_queries``。

候選表由 ``scripts/build_golden_candidates.py`` 產生。

Excel 格式（工作表名稱與欄位由候選表產生器決定）：
- 「題目」：題號、題型、題目、候選數、答案要點（請填）、漏掉的相關文件（選填）、備註
- 「候選」：題號、序、相關（1/0）、標題、藝術家、年代、來源、內容摘要、被哪些策略找到、
  external_id、doc_id
"""

from dataclasses import dataclass, field
from pathlib import Path

import psycopg

TYPE_MAP = {"事實查詢": "fact", "關係查詢": "relation", "應拒答": "refuse"}


@dataclass
class GoldenItem:
    """一題的標註結果。"""

    key: str
    query_text: str
    query_type: str
    relevant_doc_ids: set[str] = field(default_factory=set)
    missed_external_ids: list[str] = field(default_factory=list)
    gold_answer_points: list[str] = field(default_factory=list)
    unlabeled: int = 0
    invalid: list[str] = field(default_factory=list)


def _split(value: object, separators: str = ",，\n") -> list[str]:
    text = str(value or "")
    for sep in separators[1:]:
        text = text.replace(sep, separators[0])
    return [part.strip() for part in text.split(separators[0]) if part.strip()]


def read_workbook(path: str, only: set[int] | None = None) -> list[GoldenItem]:
    """讀取標註表；``only`` 指定時只取這些題號（試標用）。

    Returns:
        依題號排序的 ``GoldenItem``；未標註列數與非 0／1 的值記在 ``unlabeled``／``invalid``。
    """
    import openpyxl  # noqa: PLC0415 -- 只有匯入需要，DB 測試不必安裝

    prefix = Path(path).stem.replace("-candidates", "")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    items: dict[int, GoldenItem] = {}
    for row in wb["題目"].iter_rows(min_row=2, values_only=True):
        no, typ, text, _n, points, missed = row[:6]
        if no is None or (only and int(no) not in only):
            continue
        items[int(no)] = GoldenItem(
            key=f"{prefix}#{int(no)}",
            query_text=str(text),
            query_type=TYPE_MAP.get(str(typ), str(typ)),
            missed_external_ids=_split(missed),
            gold_answer_points=_split(points, "\n"),
        )
    for row in wb["候選"].iter_rows(min_row=2, values_only=True):
        no, seq, label, doc_id = row[0], row[1], row[2], row[10]
        if no is None or int(no) not in items:
            continue
        item = items[int(no)]
        if label in (None, ""):
            item.unlabeled += 1
        elif str(label).strip() == "1":
            item.relevant_doc_ids.add(str(doc_id))
        elif str(label).strip() != "0":
            item.invalid.append(f"第{seq}列「{label}」")
    return [items[no] for no in sorted(items)]


def problems(items: list[GoldenItem]) -> list[str]:
    """匯入前檢查：未標註、非法值、可答題沒有任何相關文件。"""
    found = []
    for item in items:
        if item.unlabeled:
            found.append(f"{item.key}：{item.unlabeled} 列未標註")
        if item.invalid:
            found.append(f"{item.key}：相關欄不是 0/1：{'、'.join(item.invalid)}")
        if item.query_type != "refuse" and not (item.relevant_doc_ids or item.missed_external_ids):
            found.append(f"{item.key}（{item.query_type}）：沒有任何相關文件")
    return found


def upsert_golden(
    conn: psycopg.Connection, items: list[GoldenItem], annotator: str
) -> dict[str, int]:
    """寫入或更新 ``eval.golden_queries``（以 ``golden_key`` 對應）。

    「漏掉的相關文件」以 external_id 查 catalog 換成 doc_id 後併入相關文件。

    Raises:
        ValueError: 「漏掉的相關文件」中有 catalog 找不到的 external_id。
    """
    stats = {"inserted": 0, "updated": 0}
    with conn.transaction():
        for item in items:
            relevant = set(item.relevant_doc_ids)
            if item.missed_external_ids:
                rows = conn.execute(
                    "SELECT external_id, doc_id::text FROM catalog.documents"
                    " WHERE external_id = ANY(%s)",
                    (item.missed_external_ids,),
                ).fetchall()
                unknown = set(item.missed_external_ids) - {r[0] for r in rows}
                if unknown:
                    raise ValueError(f"{item.key}：catalog 找不到 {sorted(unknown)}")
                relevant |= {r[1] for r in rows}
            inserted = conn.execute(
                "INSERT INTO eval.golden_queries (golden_key, query_text, query_type,"
                " relevant_doc_ids, gold_answer_points, annotator, annotated_at, status)"
                " VALUES (%s, %s, %s, %s::uuid[], %s, %s, now(), 'active')"
                " ON CONFLICT (golden_key) DO UPDATE SET query_text = EXCLUDED.query_text,"
                " query_type = EXCLUDED.query_type, relevant_doc_ids = EXCLUDED.relevant_doc_ids,"
                " gold_answer_points = EXCLUDED.gold_answer_points, annotator = EXCLUDED.annotator,"
                " annotated_at = EXCLUDED.annotated_at, status = 'active'"
                " RETURNING (xmax = 0)",
                (
                    item.key,
                    item.query_text,
                    item.query_type,
                    sorted(relevant),
                    item.gold_answer_points,
                    annotator,
                ),
            ).fetchone()[0]
            stats["inserted" if inserted else "updated"] += 1
    return stats
